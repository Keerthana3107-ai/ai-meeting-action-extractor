import io
import re
from pathlib import Path
from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {"txt", "pdf", "docx"}

def allowed_file(filename: str) -> bool:
    """Check if the uploaded file has a supported extension."""
    if not filename or "." not in filename:
        return False
    ext = filename.rsplit(".", 1)[1].lower()
    return ext in ALLOWED_EXTENSIONS

def clean_transcript(text: str) -> str:
    """Clean and normalize transcript text."""
    if not text:
        return ""
    
    # Replace non-breaking spaces and smart quotes
    text = text.replace("\xa0", " ")
    text = text.replace("“", '"').replace("”", '"')
    text = text.replace("‘", "'").replace("’", "'")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    
    # Clean multiple spaces on lines
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    
    # Merge excessive empty lines
    cleaned_lines = []
    prev_empty = False
    for line in lines:
        if not line:
            if not prev_empty:
                cleaned_lines.append("")
                prev_empty = True
        else:
            cleaned_lines.append(line)
            prev_empty = False
            
    return "\n".join(cleaned_lines).strip()

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract and decode text from TXT bytes."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")

def extract_text_from_pdf(stream_or_path) -> str:
    """Extract text from a PDF stream or file path using pypdf."""
    reader = PdfReader(stream_or_path)
    pages_text = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            pages_text.append(page_text)
    return "\n\n".join(pages_text)

def extract_text_from_docx(stream_or_path) -> str:
    """Extract text from a DOCX stream or file path using python-docx."""
    doc = Document(stream_or_path)
    paragraphs = []
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())
            
    # Also extract any text inside tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
                
    return "\n".join(paragraphs)

def extract_text_from_upload(file_storage) -> str:
    """
    Extracts text from a Flask FileStorage object.
    Supports .txt, .pdf, and .docx formats.
    """
    filename = file_storage.filename or ""
    if not allowed_file(filename):
        raise ValueError(f"Unsupported file format for '{filename}'. Allowed: {', '.join(ALLOWED_EXTENSIONS)}")
    
    ext = filename.rsplit(".", 1)[1].lower()
    file_bytes = file_storage.read()
    file_stream = io.BytesIO(file_bytes)
    
    if ext == "txt":
        raw_text = extract_text_from_txt(file_bytes)
    elif ext == "pdf":
        raw_text = extract_text_from_pdf(file_stream)
    elif ext == "docx":
        raw_text = extract_text_from_docx(file_stream)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
        
    return clean_transcript(raw_text)
